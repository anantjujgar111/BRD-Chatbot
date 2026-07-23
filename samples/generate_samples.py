"""Generate sample Excel/CSV files for BRD chatbot testing."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

OUTPUT_DIR = Path(__file__).resolve().parent / "excel"


def create_structured_functional_requirements() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Functional Requirements"

    headers = ["Req ID", "Module", "Requirement", "Priority", "Owner", "Status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    rows = [
        ["FR-001", "Authentication", "System shall support email and password login", "High", "Product", "Approved"],
        ["FR-002", "Authentication", "System shall support SSO via Azure AD", "High", "Security", "Approved"],
        ["FR-003", "Payments", "System shall process card payments with 3DS", "High", "Payments", "In Review"],
        ["FR-004", "Payments", "System shall support partial and full refunds", "Medium", "Payments", "Approved"],
        ["FR-005", "Reporting", "Admin shall export transaction reports to Excel", "Medium", "Ops", "Draft"],
        ["FR-006", "Notifications", "System shall send email alerts for failed payments", "Low", "Product", "Draft"],
    ]
    for row in rows:
        ws.append(row)

    wb.save(OUTPUT_DIR / "01_structured_functional_requirements.xlsx")


def create_structured_stakeholders() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stakeholders"

    headers = ["Stakeholder", "Role", "Department", "Interest", "Influence", "Contact"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    rows = [
        ["Ravi Mehta", "Product Owner", "Business", "Delivery timeline", "High", "ravi@example.com"],
        ["Sneha Patil", "Compliance Lead", "Legal", "Regulatory alignment", "High", "sneha@example.com"],
        ["Arjun Kulkarni", "Tech Lead", "Engineering", "Architecture stability", "High", "arjun@example.com"],
        ["Finance Ops", "Approver", "Finance", "Settlement accuracy", "Medium", "finance@example.com"],
        ["Customer Support", "End User Rep", "Support", "Ease of use", "Medium", "support@example.com"],
    ]
    for row in rows:
        ws.append(row)

    wb.save(OUTPUT_DIR / "02_structured_stakeholders_matrix.xlsx")


def create_structured_nfr_csv() -> None:
    content = """Category,Requirement,Target,Verification Method
Performance,API response time for payment status,<= 2 seconds,Load test
Security,All PII must be encrypted at rest,AES-256,Security audit
Availability,Core payment service uptime,99.9% monthly,Monitoring dashboard
Scalability,Support 10x peak traffic during sales events,Auto scaling,Stress test
Audit,All admin actions must be logged,100% coverage,Log review
"""
    (OUTPUT_DIR / "03_structured_non_functional_requirements.csv").write_text(content, encoding="utf-8")


def create_unstructured_client_notes() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Client Notes"

    ws.merge_cells("A1:F1")
    ws["A1"] = "Workshop Notes - Not Final"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Random discussion points"
    ws["A5"] = "Customer wants faster refunds - maybe same day?"
    ws["B5"] = "check with bank"
    ws["A6"] = "Login issue for dealers in rural areas - OTP delays"
    ws["A8"] = "Someone mentioned GST invoice format change in Q3"
    ws.merge_cells("C8:E8")
    ws["C8"] = "Need legal review before committing"

    ws["A11"] = "Scope confusion"
    ws["A12"] = "Mobile app NOT in phase 1"
    ws["B12"] = "but management keeps asking"
    ws["A14"] = "Reporting"
    ws["A15"] = "daily settlement file by 9am"
    ws["C15"] = "format unknown"

    ws2 = wb.create_sheet("Call Log")
    ws2["B2"] = "Date"
    ws2["D2"] = "Notes"
    ws2["B4"] = "12-Jul"
    ws2.merge_cells("D4:G6")
    ws2["D4"] = "Client said: 'we cannot change excel format, teams send files like this only'"

    wb.save(OUTPUT_DIR / "04_unstructured_client_notes.xlsx")


def create_unstructured_workshop_capture() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Workshop Day 1"

    ws["C3"] = "Feature"
    ws["E3"] = "Details"
    ws["C4"] = "KYC upload"
    ws["E4"] = "PAN + Aadhaar mandatory"
    ws["C7"] = "Assumption: all users have email"
    ws["A10"] = "OUT OF SCOPE -> WhatsApp bot"
    ws.merge_cells("A12:C12")
    ws["A12"] = "Process sketch (rough)"
    ws["A14"] = "1) user registers"
    ws["A15"] = "2) ops verifies docs"
    ws["A16"] = "3) account activated"
    ws["E16"] = "SLA 24h? TBD"

    ws["G5"] = "Risk"
    ws["G6"] = "Third-party KYC vendor downtime"
    ws["G8"] = "Dependency on legacy CRM export"

    wb.save(OUTPUT_DIR / "05_unstructured_workshop_capture.xlsx")


def create_unstructured_legacy_export() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export"

    ws["A1"] = "Legacy Requirements Dump"
    ws["A4"] = "Req#"
    ws["C4"] = "Description"
    ws["F4"] = "Priority"

    ws["A6"] = "R-101"
    ws.merge_cells("C6:E7")
    ws["C6"] = "Dealer portal must allow bulk order upload via excel"

    ws["A9"] = "R-102"
    ws["C9"] = "Inventory sync every 15 min"
    ws["F9"] = "High"

    ws["A12"] = "misc notes"
    ws["C12"] = "password reset link expires in 30 min"
    ws["C13"] = "audit trail required for price changes"

    ws2 = wb.create_sheet("Sheet2")
    ws2["D2"] = "Acceptance Criteria"
    ws2["D4"] = "Refund request visible in admin within 5 minutes"
    ws2["D5"] = "Failed payment retry up to 3 times"

    wb.save(OUTPUT_DIR / "06_unstructured_legacy_export.xlsx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    create_structured_functional_requirements()
    create_structured_stakeholders()
    create_structured_nfr_csv()
    create_unstructured_client_notes()
    create_unstructured_workshop_capture()
    create_unstructured_legacy_export()
    print(f"Created sample files in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
