from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd


@dataclass
class ExtractedBlock:
    file_name: str
    sheet_name: str
    cell_range: str
    text: str
    block_type: str = "unknown"


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _extract_merged_map(sheet) -> dict[tuple[int, int], str]:
    merged_values: dict[tuple[int, int], str] = {}
    for merged_range in sheet.merged_cells.ranges:
        top_left = sheet.cell(merged_range.min_row, merged_range.min_col).value
        text = _clean(top_left)
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_values[(row, col)] = text
    return merged_values


def _cell_text(sheet, row: int, col: int, merged_map: dict[tuple[int, int], str]) -> str:
    if (row, col) in merged_map:
        return merged_map[(row, col)]
    return _clean(sheet.cell(row, col).value)


def _infer_block_type(lines: list[str]) -> str:
    joined = " ".join(lines).lower()
    if any(word in joined for word in ["requirement", "shall", "must", "should"]):
        return "requirement"
    if any(word in joined for word in ["assumption", "constraint", "dependency"]):
        return "assumption"
    if any(word in joined for word in ["process", "workflow", "step"]):
        return "process"
    if any(word in joined for word in ["stakeholder", "actor", "user role"]):
        return "stakeholder"
    return "general"


def extract_excel_blocks(file_path: Path, file_name: str) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    suffix = file_path.suffix.lower()

    if suffix == ".csv":
        return _extract_csv_blocks(file_path, file_name)

    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=False)
    except Exception:
        return _extract_with_pandas(file_path, file_name)

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        merged_map = _extract_merged_map(sheet)
        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0

        if max_row == 0 or max_col == 0:
            continue

        current_lines: list[str] = []
        start_row = 1
        start_col = 1
        last_row = 1
        last_col = 1

        def flush_block(end_row: int, end_col: int) -> None:
            nonlocal current_lines, start_row, start_col
            cleaned = [line for line in current_lines if line]
            if not cleaned:
                current_lines = []
                return
            cell_range = f"{openpyxl.utils.get_column_letter(start_col)}{start_row}:{openpyxl.utils.get_column_letter(end_col)}{end_row}"
            blocks.append(
                ExtractedBlock(
                    file_name=file_name,
                    sheet_name=sheet_name,
                    cell_range=cell_range,
                    text="\n".join(cleaned),
                    block_type=_infer_block_type(cleaned),
                )
            )
            current_lines = []

        for row in range(1, max_row + 1):
            row_values: list[str] = []
            for col in range(1, max_col + 1):
                text = _cell_text(sheet, row, col, merged_map)
                if text:
                    row_values.append(text)

            if row_values:
                if not current_lines:
                    start_row = row
                    start_col = 1
                current_lines.append(" | ".join(row_values))
                last_row = row
                last_col = max_col
            elif current_lines:
                flush_block(last_row, last_col)

        if current_lines:
            flush_block(last_row, last_col)

    workbook.close()
    return blocks


def _extract_csv_blocks(file_path: Path, file_name: str) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    try:
        df = pd.read_csv(file_path, dtype=str, keep_default_na=False)
    except Exception:
        return blocks

    for idx, row in df.iterrows():
        values = [_clean(v) for v in row.tolist() if _clean(v)]
        if not values:
            continue
        blocks.append(
            ExtractedBlock(
                file_name=file_name,
                sheet_name="CSV",
                cell_range=f"row:{int(idx) + 2}",
                text=" | ".join(values),
                block_type=_infer_block_type(values),
            )
        )
    return blocks


def _extract_with_pandas(file_path: Path, file_name: str) -> list[ExtractedBlock]:
    blocks: list[ExtractedBlock] = []
    try:
        sheets = pd.read_excel(file_path, sheet_name=None, dtype=str, keep_default_na=False)
    except Exception:
        return blocks

    for sheet_name, df in sheets.items():
        for idx, row in df.iterrows():
            values = [_clean(v) for v in row.tolist() if _clean(v)]
            if not values:
                continue
            blocks.append(
                ExtractedBlock(
                    file_name=file_name,
                    sheet_name=str(sheet_name),
                    cell_range=f"row:{int(idx) + 2}",
                    text=" | ".join(values),
                    block_type=_infer_block_type(values),
                )
            )
    return blocks
